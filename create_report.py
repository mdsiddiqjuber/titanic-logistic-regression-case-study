from pathlib import Path
import json
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parent
FIG = BASE / 'figures'
OUT = BASE / 'outputs'
metrics = json.loads((OUT/'metrics.json').read_text())
findings = json.loads((OUT/'statistical_findings.json').read_text())
coefs = pd.read_csv(OUT/'model_coefficients.csv')
df = pd.read_csv(BASE/'titanic.csv')

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.68)
sec.bottom_margin = Inches(0.68)
sec.left_margin = Inches(0.78)
sec.right_margin = Inches(0.78)

# Palette
NAVY = '17365D'
BLUE = '2F75B5'
TEAL = '0F6B78'
PALE = 'EAF2F8'
LIGHT = 'F3F6F9'
GRAY = '5B6573'
WHITE = 'FFFFFF'

# Typography
styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal'].font.size = Pt(10)
styles['Normal'].font.color.rgb = RGBColor(38, 48, 61)
styles['Normal'].paragraph_format.space_after = Pt(5)
styles['Normal'].paragraph_format.line_spacing = 1.08
for stylename, size, color in [('Title', 29, NAVY), ('Heading 1', 19, NAVY), ('Heading 2', 13, TEAL), ('Heading 3', 11, BLUE)]:
    st = styles[stylename]
    st.font.name = 'Aptos Display' if stylename in ('Title','Heading 1') else 'Aptos'
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(7)
    st.paragraph_format.space_after = Pt(5)

# Header/footer
header = sec.header.paragraphs[0]
header.text = 'FUNDAMENTALS OF DATA SCIENCE  |  CASE STUDY'
header.style = styles['Caption']
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for r in header.runs:
    r.font.name='Aptos'; r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(GRAY)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run('Titanic Survival Prediction with Logistic Regression  •  Page ')
run.font.size=Pt(8); run.font.color.rgb=RGBColor.from_string(GRAY)
fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'),'begin')
instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'),'preserve'); instrText.text=' PAGE '
fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'),'end')
run._r.append(fldChar1); run._r.append(instrText); run._r.append(fldChar2)

# Helpers
def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def set_cell_margins(cell, top=70, start=100, bottom=70, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m, v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = tcMar.find(qn('w:'+m))
        if node is None:
            node = OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'),'dxa')

def format_table(table, header=True, font_size=9):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.name='Aptos'; r.font.size=Pt(font_size)
            if header and i == 0:
                shade(cell, NAVY)
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255)
            elif i % 2 == 0:
                shade(cell, LIGHT)

def add_table(headers, rows, widths=None, font_size=9):
    t = doc.add_table(rows=1, cols=len(headers))
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = str(h)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Inches(w)
    format_table(t, True, font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

def add_bullet(text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(2)
    p.add_run(text)
    return p

def add_caption(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text); r.italic = True; r.font.size = Pt(8.5); r.font.color.rgb=RGBColor.from_string(GRAY)

def add_figure(filename, caption, width=5.95):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(FIG/filename), width=Inches(width))
    add_caption(caption)

def add_interpretation(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run('Interpretation. '); r.bold = True; r.font.color.rgb=RGBColor.from_string(TEAL)
    p.add_run(text)

def add_callout(text):
    t = doc.add_table(rows=1, cols=1)
    c = t.cell(0,0); c.text = text; shade(c, PALE); set_cell_margins(c,110,150,110,150)
    for p in c.paragraphs:
        p.paragraph_format.space_after=Pt(0)
        for r in p.runs:
            r.font.size=Pt(9.5); r.font.color.rgb=RGBColor.from_string(NAVY)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

_pagebreak_count = 0
def pagebreak():
    global _pagebreak_count
    _pagebreak_count += 1
    if _pagebreak_count == 1:  # Keep the cover page separate; allow report content to flow naturally.
        doc.add_page_break()

# PAGE 1: Cover
p = doc.add_paragraph(); p.paragraph_format.space_before=Pt(52); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('CASE STUDY REPORT'); r.bold=True; r.font.size=Pt(14); r.font.color.rgb=RGBColor.from_string(TEAL)
p = doc.add_paragraph(); p.paragraph_format.space_before=Pt(18); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('TITANIC SURVIVAL\nPREDICTION'); r.bold=True; r.font.name='Aptos Display'; r.font.size=Pt(31); r.font.color.rgb=RGBColor.from_string(NAVY)
p = doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Using Logistic Regression'); r.font.size=Pt(18); r.font.color.rgb=RGBColor.from_string(BLUE)
p = doc.add_paragraph(); p.paragraph_format.space_before=Pt(18); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Exploratory Data Analysis • Preprocessing • Model Evaluation • Interpretation'); r.font.size=Pt(10); r.font.color.rgb=RGBColor.from_string(GRAY)
for _ in range(2): doc.add_paragraph()
add_table(['Academic detail','Information'], [
    ['Name','Md. Siddiq Juber'],
    ['Roll Number','[ENTER ROLL NUMBER]'],
    ['Section','[ENTER SECTION]'],
    ['Year','III Year'],
    ['Subject','Fundamentals of Data Science'],
    ['Faculty Name','[ENTER FACULTY NAME]'],
], widths=[2.0,4.0], font_size=10)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(18)
r=p.add_run('Dataset: Titanic passenger records (891 rows)'); r.italic=True; r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string(GRAY)
pagebreak()

# PAGE 2: Summary and problem
doc.add_heading('Executive Summary', level=1)
doc.add_paragraph(
    'This case study develops a binary classification model to estimate whether a passenger in the supplied Titanic dataset survived. '
    'The analysis includes descriptive exploration, visualizations, missing-value handling, categorical encoding, logistic-regression training, '
    'held-out evaluation, coefficient interpretation, and a saved model pipeline for a Streamlit demo.'
)
add_callout(
    f"Main result: on a stratified 20% test set ({metrics['testing_rows']} passengers), the model achieved "
    f"{metrics['accuracy']:.1%} accuracy and ROC-AUC {metrics['roc_auc']:.3f}. These figures describe this fixed split and are not a guarantee of future performance."
)
doc.add_heading('1. Problem Statement', level=1)
doc.add_paragraph(
    'Given passenger characteristics, predict the binary target Survived (0 = did not survive, 1 = survived). '
    'The goal is to apply the data-science workflow prescribed in the assignment and explain the findings rather than present charts or scores without interpretation.'
)
doc.add_heading('Objectives', level=2)
for txt in [
    'Inspect feature types, descriptive statistics, class balance, missing values, distributions, and relationships.',
    'Prepare usable predictors by imputing missing values, scaling numeric inputs, and encoding categorical inputs.',
    'Train logistic regression using a reproducible, stratified train/test split.',
    'Report accuracy, precision, recall, F1-score, ROC-AUC, a ROC curve, and a confusion matrix.',
    'Interpret model coefficients and make the trained pipeline available to a Streamlit app.'
]: add_bullet(txt)
doc.add_heading('2. Dataset Overview', level=1)
add_table(['Item','Finding'], [
    ['Rows / columns', f"{df.shape[0]} rows and {df.shape[1]} columns"],
    ['Target', 'Survived: 0 = did not survive; 1 = survived'],
    ['Survived', f"{int(df['Survived'].sum())} ({df['Survived'].mean():.1%})"],
    ['Did not survive', f"{int((df['Survived']==0).sum())} ({(df['Survived']==0).mean():.1%})"],
    ['Numeric predictors used', 'Age, SibSp, Parch, Fare'],
    ['Categorical predictors used', 'Pclass, Sex, Embarked'],
], widths=[1.9,4.3], font_size=9)
doc.add_paragraph(
    'PassengerId, Name, Ticket, and Cabin are not included in this baseline. PassengerId is an identifier; Name and Ticket are high-cardinality text fields that could require feature engineering; Cabin is missing for most records.'
)
pagebreak()

# PAGE 3: target distribution
doc.add_heading('3. Exploratory Data Analysis', level=1)
doc.add_heading('3.1 Target class distribution', level=2)
add_figure('01_survival_distribution.png', 'Figure 1. Count and proportion of passengers in each survival class.', width=5.8)
add_interpretation(
    f"The data contains {int((df['Survived']==0).sum())} non-survivors and {int(df['Survived'].sum())} survivors. "
    f"The observed survival rate is {df['Survived'].mean():.1%}; non-survivors form the larger class. Because the target is not evenly split, the report uses precision, recall, F1-score, and ROC-AUC in addition to accuracy."
)
doc.add_heading('Observed target values', level=2)
add_table(['Outcome','Count','Share'], [
    ['Did not survive (0)', f"{int((df['Survived']==0).sum())}", f"{(df['Survived']==0).mean():.1%}"],
    ['Survived (1)', f"{int(df['Survived'].sum())}", f"{df['Survived'].mean():.1%}"],
], widths=[2.8,1.4,1.4], font_size=9)
pagebreak()

# PAGE 4: survival by sex
doc.add_heading('3.2 Observed survival rate by sex', level=2)
add_figure('02_survival_by_sex.png', 'Figure 2. Percentage of passengers who survived, grouped by recorded sex.', width=5.8)
add_interpretation(
    f"The observed survival rate was {findings['survival_rate_by_sex']['female']:.1%} among female passengers and "
    f"{findings['survival_rate_by_sex']['male']:.1%} among male passengers. This is a large descriptive difference in the supplied data and is reflected in the fitted model, but it should not be interpreted as a causal effect."
)
doc.add_heading('3.3 Observed survival rate by passenger class', level=2)
add_figure('03_survival_by_class.png', 'Figure 3. Percentage of passengers who survived, grouped by passenger class.', width=5.8)
add_interpretation(
    f"Observed survival rates were {findings['survival_rate_by_class']['1']:.1%} in first class, "
    f"{findings['survival_rate_by_class']['2']:.1%} in second class, and {findings['survival_rate_by_class']['3']:.1%} in third class. "
    'Survival rates decrease across these class groups in the dataset. Class may be associated with several other passenger characteristics, so this pattern alone does not explain the mechanism.'
)
pagebreak()

# PAGE 5: age and missingness
doc.add_heading('3.4 Age distribution by survival outcome', level=2)
add_figure('04_age_distribution_by_survival.png', 'Figure 4. Age distributions for the two observed survival outcomes.', width=5.15)
add_interpretation(
    'The age distributions overlap, meaning age alone cannot separate survivors from non-survivors. The plot includes available ages; missing ages are handled later inside the training pipeline so that test data do not determine imputation values.'
)
doc.add_heading('3.5 Missing values', level=2)
add_figure('05_missing_values.png', 'Figure 5. Percentage of missing observations in columns with at least one missing value.', width=5.15)
add_interpretation(
    f"Age is missing in {findings['missing_age_count']} rows ({df['Age'].isna().mean():.1%}), Cabin in "
    f"{findings['missing_cabin_count']} rows ({df['Cabin'].isna().mean():.1%}), and Embarked in {findings['missing_embarked_count']} rows ({df['Embarked'].isna().mean():.1%}). "
    'Cabin is omitted from the baseline because of its high missingness; the selected predictors are imputed by the model pipeline.'
)
pagebreak()

# PAGE 6: correlation
doc.add_heading('3.6 Correlation among selected numeric columns', level=2)
add_figure('06_correlation_heatmap.png', 'Figure 6. Pairwise Pearson correlation coefficients among selected numeric columns.', width=5.8)
add_interpretation(
    'The heatmap summarizes linear pairwise correlation. In this dataset, Survived has a negative correlation with Pclass as its numeric coding increases from first to third class, and a positive correlation with Fare. This is a descriptive comparison; Pearson correlation does not establish causation, and Pclass is treated as categorical in the model rather than assuming equal numeric spacing.'
)
doc.add_heading('Feature-type review', level=2)
add_table(['Feature type','Columns','Model treatment'], [
    ['Integer / numeric', 'PassengerId, Survived, Pclass, SibSp, Parch', 'Survived is target; PassengerId excluded; Pclass encoded categorically; SibSp/Parch numeric'],
    ['Floating-point', 'Age, Fare', 'Median imputation and standardization'],
    ['Categorical / text', 'Name, Sex, Ticket, Cabin, Embarked', 'Sex and Embarked encoded; Name/Ticket omitted; Cabin omitted for missingness'],
], widths=[1.4,2.3,3.0], font_size=8)
pagebreak()

# PAGE 7 methodology
doc.add_page_break()
doc.add_heading('4. Methodology and Data Preprocessing', level=1)
doc.add_heading('4.1 Predictor selection', level=2)
doc.add_paragraph(
    'The target is Survived. Seven predictors are retained: Age, SibSp, Parch, Fare, Pclass, Sex, and Embarked. '
    'The identifier PassengerId is excluded; Name and Ticket are left out because this baseline does not engineer information from high-cardinality text; Cabin is excluded because 77.1% of its values are missing.'
)
doc.add_heading('4.2 Train/test split', level=2)
add_table(['Setting','Choice'], [
    ['Training fraction','80% (712 rows)'],
    ['Test fraction','20% (179 rows)'],
    ['Random seed','42'],
    ['Stratification','Yes, preserving the target class proportion in both sets'],
    ['Classification threshold','0.50 probability'],
], widths=[2.0,4.2], font_size=9)
doc.add_heading('4.3 Preprocessing pipeline', level=2)
for txt in [
    'Numeric features (Age, SibSp, Parch, Fare): median imputation followed by standard scaling.',
    'Categorical features (Pclass, Sex, Embarked): most-frequent imputation followed by one-hot encoding with the first category dropped as a reference.',
    'All preprocessing is fit on the training split as part of a scikit-learn Pipeline and ColumnTransformer, then applied to test data using the learned transformations.',
    'Classifier: scikit-learn LogisticRegression with max_iter = 2000 and random_state = 42.'
]: add_bullet(txt)
doc.add_heading('Why logistic regression?', level=2)
doc.add_paragraph(
    'The target is binary and logistic regression estimates the probability of the positive class. It is a transparent baseline because coefficients can be inspected, but its results depend on the selected features, preprocessing, split, and classification threshold.'
)
pagebreak()

# PAGE 8: metrics and ROC
doc.add_page_break()
doc.add_heading('5. Model Evaluation', level=1)
doc.add_paragraph('The model is evaluated once on the held-out test set of 179 passengers. All predictions in the following table use the default threshold of 0.50.')
metric_rows = [
    ['Accuracy', f"{metrics['accuracy']:.4f}", f"{metrics['accuracy']:.1%}", 'Fraction of all test predictions that are correct.'],
    ['Precision', f"{metrics['precision']:.4f}", f"{metrics['precision']:.1%}", 'Among passengers predicted to survive, the fraction who survived.'],
    ['Recall', f"{metrics['recall']:.4f}", f"{metrics['recall']:.1%}", 'Among actual survivors, the fraction identified by the model.'],
    ['F1-score', f"{metrics['f1_score']:.4f}", f"{metrics['f1_score']:.1%}", 'Harmonic mean of precision and recall.'],
    ['ROC-AUC', f"{metrics['roc_auc']:.4f}", f"{metrics['roc_auc']:.3f}", 'Ranking ability over a range of classification thresholds.'],
]
add_table(['Metric','Score','Percent / value','Meaning'], metric_rows, widths=[1.05,0.9,1.15,3.2], font_size=8)
doc.add_heading('5.1 ROC curve', level=2)
add_figure('07_roc_curve.png', 'Figure 7. True-positive rate versus false-positive rate across decision thresholds.', width=5.8)
add_interpretation(
    f"The ROC-AUC is {metrics['roc_auc']:.3f}, above the 0.500 random-classifier reference. It indicates the model ranks positive cases above negative cases reasonably well on this split. ROC-AUC is not the same as accuracy or the probability that any individual prediction is correct."
)
pagebreak()

# PAGE 9 confusion matrix
doc.add_heading('5.2 Confusion matrix and error analysis', level=2)
add_figure('08_confusion_matrix.png', 'Figure 8. Counts of correct and incorrect predictions on the test set.', width=5.5)
cm = metrics['confusion_matrix']
tn, fp = cm[0]
fn, tp = cm[1]
add_table(['Outcome','Count','Meaning'], [
    ['True negative (TN)', tn, 'Predicted did not survive; actual did not survive.'],
    ['False positive (FP)', fp, 'Predicted survived; actual did not survive.'],
    ['False negative (FN)', fn, 'Predicted did not survive; actual survived.'],
    ['True positive (TP)', tp, 'Predicted survived; actual survived.'],
], widths=[1.8,0.8,3.8], font_size=9)
add_interpretation(
    f"The model correctly classified {tn+tp} of {metrics['testing_rows']} passengers. It produced {fp} false positives and missed {fn} actual survivors (false negatives). "
    'Recall is lower than precision here because the model misses about one-third of the actual survivors at the 0.50 threshold. A different threshold may raise recall but can also increase false positives.'
)
doc.add_heading('Metric interpretation at a glance', level=2)
add_bullet(f"Accuracy ({metrics['accuracy']:.1%}) gives the overall correctness rate, but can hide differences between classes.")
add_bullet(f"Precision ({metrics['precision']:.1%}) is useful when false positive survivor predictions matter.")
add_bullet(f"Recall ({metrics['recall']:.1%}) shows that about two-thirds of actual survivors are identified at the current threshold.")
add_bullet(f"F1 ({metrics['f1_score']:.1%}) summarizes the precision/recall balance; ROC-AUC ({metrics['roc_auc']:.3f}) summarizes ranking across thresholds.")
pagebreak()

# PAGE 10 coefficients
doc.add_heading('6. Interpretation of Logistic Regression Coefficients', level=1)
doc.add_paragraph(
    'A logistic-regression coefficient is on the log-odds scale. Positive coefficients increase modeled survival log-odds and negative coefficients decrease them, holding the other included features constant. The exponentiated coefficient is the odds ratio.'
)
add_figure('09_model_coefficients.png', 'Figure 9. Coefficients for standardized numeric and one-hot encoded categorical features.', width=5.85)
# Helper for getting one coefficient by exact feature id
coef_index = coefs.set_index('Feature')
key_names = ['cat__Sex_male','cat__Pclass_3','cat__Pclass_2','num__Age','num__Fare','num__SibSp','cat__Embarked_S','cat__Embarked_Q']
rows=[]
labels={
 'cat__Sex_male':'Male vs female (reference)',
 'cat__Pclass_3':'Third vs first class (reference)',
 'cat__Pclass_2':'Second vs first class (reference)',
 'num__Age':'Age (per 1 SD increase)',
 'num__Fare':'Fare (per 1 SD increase)',
 'num__SibSp':'Siblings/spouses (per 1 SD increase)',
 'cat__Embarked_S':'Southampton vs Cherbourg (reference)',
 'cat__Embarked_Q':'Queenstown vs Cherbourg (reference)',
}
for key in key_names:
    if key in coef_index.index:
        row=coef_index.loc[key]
        rows.append([labels[key], f"{row['Coefficient']:.3f}", f"{row['Odds ratio']:.3f}"])
add_table(['Feature comparison','Coefficient','Odds ratio'], rows, widths=[3.5,1.2,1.2], font_size=8.5)
add_interpretation(
    'The male indicator has a strongly negative coefficient (odds ratio about 0.08 relative to female, adjusted for the selected predictors). Third class also has a strongly negative coefficient relative to first class (odds ratio about 0.15). Age has a negative coefficient per standard deviation, while Fare has a modest positive coefficient. These are model associations, not causal effects. Numeric coefficients are interpreted per one standard deviation because the numeric columns were scaled. The reference categories are first class, female, and Cherbourg.'
)
add_callout('The coefficient plot does not report p-values or confidence intervals. Do not call a feature statistically significant based only on its coefficient size; formal statistical inference was outside the scope of this baseline analysis.')
pagebreak()

# PAGE 11 conclusion/deployment/repo
doc.add_heading('7. Conclusion', level=1)
doc.add_paragraph(
    f"The case study completes an end-to-end binary-classification workflow on 891 Titanic passenger records. The model achieved {metrics['accuracy']:.1%} accuracy, "
    f"{metrics['precision']:.1%} precision, {metrics['recall']:.1%} recall, an F1-score of {metrics['f1_score']:.1%}, and ROC-AUC {metrics['roc_auc']:.3f} on the held-out test split. "
    'EDA showed substantial observed survival-rate differences by recorded sex and passenger class. The model coefficients show similar associations after adjustment for the selected predictors.'
)
doc.add_heading('Limitations and future improvements', level=2)
for txt in [
    'Results depend on a single stratified split; cross-validation would give a more stable performance estimate.',
    'The dataset is historical and may reflect social and sampling biases; predictions should not be generalized casually.',
    'Cabin is mostly missing, and Name/Ticket were not feature-engineered in this baseline.',
    'The 0.50 threshold is a default; threshold tuning, probability calibration, and a cost-sensitive error analysis could be explored.',
    'Coefficient sizes are not proof of causality; confidence intervals and statistical tests were not calculated.'
]: add_bullet(txt)
doc.add_heading('8. Streamlit Deployment', level=1)
doc.add_paragraph(
    'The project includes app.py. It loads the complete preprocessing + logistic-regression pipeline and accepts passenger class, sex, age, siblings/spouses, parents/children, fare, and embarkation port. It displays estimated survival probability and a class prediction at the 0.50 threshold.'
)
add_callout('Run locally: install packages with `pip install -r requirements.txt`, run `python analyze.py` to create the trained model, then run `streamlit run app.py`.')
doc.add_heading('9. GitHub Submission', level=1)
doc.add_paragraph(
    'The repository package includes the executed Jupyter notebook, Python analysis and Streamlit files, requirements, dataset, figure files, model artifacts, metric outputs, and README. Create a GitHub repository, then initialize, commit, add the remote, and push the project as described in README.md. The GitHub URL should be submitted alongside this PDF.'
)
doc.add_paragraph(
    'Before final submission, replace the placeholders on the title page with the correct roll number, section, and faculty name, and verify the repository visibility/access rules required by the course.'
)

# Core properties
doc.core_properties.title = 'Titanic Survival Prediction Using Logistic Regression'
doc.core_properties.subject = 'Fundamentals of Data Science case study'
doc.core_properties.author = 'Md. Siddiq Juber'
doc.core_properties.keywords = 'Titanic, logistic regression, data science, classification, EDA'

out = BASE / 'Titanic_Logistic_Regression_Case_Study.docx'
doc.save(out)
print(out)
